import subprocess
import os
import shutil


def find_libreoffice():
    """
    Locate the LibreOffice command-line binary. Checks PATH first, then the
    typical Windows install locations, similar to how ocr_engine.py falls
    back to a known path for Tesseract.
    """
    for name in ("soffice", "soffice.exe"):
        found = shutil.which(name)
        if found:
            return found

    windows_candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for path in windows_candidates:
        if os.path.exists(path):
            return path

    return None


def convert_with_libreoffice(input_path, output_dir, target_format):
    """
    Converts input_path to target_format (e.g. "pdf", "docx", "pptx", "xlsx")
    using LibreOffice's headless command-line conversion. Returns the path
    to the converted file, or raises RuntimeError with a clear message if
    LibreOffice isn't installed/found.
    """
    soffice = find_libreoffice()
    if not soffice:
        raise RuntimeError(
            "LibreOffice was not found. Please install it from "
            "https://www.libreoffice.org/download/download/ and make sure "
            "it's on your PATH, or installed at the default location."
        )

    os.makedirs(output_dir, exist_ok=True)

    result = subprocess.run(
        [
            soffice,
            "--headless",
            "--norestore",
            "--convert-to", target_format,
            "--outdir", output_dir,
            input_path,
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )

    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice conversion failed: {result.stderr or result.stdout}")

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    expected_output = os.path.join(output_dir, f"{base_name}.{target_format}")

    if not os.path.exists(expected_output):
        raise RuntimeError(
            f"LibreOffice reported success but no output file was found at {expected_output}."
        )

    return expected_output


def convert_pdf_to_pdfa(input_path, output_dir):
    """
    Converts a PDF to PDF/A using LibreOffice's PDF/A export filter. This
    re-imports the PDF into LibreOffice and re-exports it, so complex
    layouts may shift slightly - this is a best-effort conversion, not a
    guarantee of full PDF/A compliance for every document.
    """
    soffice = find_libreoffice()
    if not soffice:
        raise RuntimeError(
            "LibreOffice was not found. Please install it from "
            "https://www.libreoffice.org/download/download/ and make sure "
            "it's on your PATH, or installed at the default location."
        )

    os.makedirs(output_dir, exist_ok=True)

    result = subprocess.run(
        [
            soffice,
            "--headless",
            "--norestore",
            "--convert-to", "pdf:writer_pdf_Export:SelectPdfVersion=1",
            "--outdir", output_dir,
            input_path,
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )

    if result.returncode != 0:
        raise RuntimeError(f"PDF/A conversion failed: {result.stderr or result.stdout}")

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    expected_output = os.path.join(output_dir, f"{base_name}.pdf")

    if not os.path.exists(expected_output):
        raise RuntimeError("LibreOffice reported success but no PDF/A output file was found.")

    return expected_output


def convert_pdf_to_word(input_path, output_path):
    """
    Converts a PDF to a .docx file using pdf2docx, which is purpose-built
    for this and generally preserves layout/text better than routing
    through LibreOffice's PDF import filter.
    """
    from pdf2docx import Converter

    cv = Converter(input_path)
    cv.convert(output_path)
    cv.close()
    return output_path


def convert_pdf_to_excel(input_path, output_path):
    """
    Extracts tables from a PDF using pdfplumber and writes each page's
    table(s) into separate sheets in an .xlsx file. Works well for PDFs
    with real tabular data; won't find tables in scanned/image-only PDFs
    (those would need OCR first).
    """
    import pdfplumber
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet
    found_any_table = False

    with pdfplumber.open(input_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            for table_num, table in enumerate(tables, start=1):
                found_any_table = True
                sheet_name = f"Page{page_num}_Table{table_num}"[:31]
                ws = wb.create_sheet(title=sheet_name)
                for row in table:
                    ws.append([cell if cell is not None else "" for cell in row])

    if not found_any_table:
        ws = wb.create_sheet(title="Extracted Text")
        with pdfplumber.open(input_path) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                for line in text.split("\n"):
                    ws.append([line])

    wb.save(output_path)
    return output_path
